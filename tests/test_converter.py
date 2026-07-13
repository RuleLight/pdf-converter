from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document

from converter import ConversionOptions, convert_pdf, find_tesseract, parse_page_range, validate_docx


def make_digital_pdf(path: Path) -> None:
    pdf = fitz.open()
    first = pdf.new_page(width=595, height=842)
    first.insert_text((72, 90), "Cross-platform PDF to Word", fontsize=20)
    first.insert_text((72, 125), "Editable text, a simple table, and two pages.", fontsize=11)
    for row in range(3):
        y = 170 + row * 30
        first.draw_line((72, y), (420, y))
    for x in (72, 246, 420):
        first.draw_line((x, 170), (x, 230))
    first.insert_text((82, 192), "Name", fontsize=10)
    first.insert_text((256, 192), "Value", fontsize=10)
    first.insert_text((82, 222), "Example", fontsize=10)
    first.insert_text((256, 222), "42", fontsize=10)
    second = pdf.new_page(width=842, height=595)
    second.insert_text((72, 90), "Landscape second page", fontsize=20)
    pdf.save(path)
    pdf.close()


def make_scanned_pdf(path: Path) -> None:
    image = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 72)
    except OSError:
        font = ImageFont.load_default(size=72)
    draw.text((80, 100), "SCAN TEST 2026", fill="black", font=font)
    png_path = path.with_suffix(".png")
    image.save(png_path)
    pdf = fitz.open()
    page = pdf.new_page(width=600, height=350)
    page.insert_image(page.rect, filename=str(png_path))
    pdf.save(path)
    pdf.close()
    png_path.unlink()


class ConverterTests(unittest.TestCase):
    def test_page_range(self) -> None:
        self.assertEqual(parse_page_range("1-3, 5", 6), [0, 1, 2, 4])
        self.assertEqual(parse_page_range("全部", 2), [0, 1])

    def test_layout_and_exact(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "digital.pdf"
            make_digital_pdf(source)

            layout = convert_pdf(source, root / "layout.docx", ConversionOptions(mode="layout"))
            validate_docx(layout)
            layout_text = "\n".join(p.text for p in Document(layout).paragraphs)
            self.assertIn("Cross-platform PDF to Word", layout_text)

            exact = convert_pdf(source, root / "exact.docx", ConversionOptions(mode="exact"))
            validate_docx(exact)
            exact_doc = Document(exact)
            self.assertEqual(len(exact_doc.inline_shapes), 2)
            self.assertEqual(len(exact_doc.sections), 2)

    @unittest.skipUnless(find_tesseract(), "Tesseract is not installed")
    def test_ocr(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "scan.pdf"
            make_scanned_pdf(source)
            output = convert_pdf(
                source,
                root / "ocr.docx",
                ConversionOptions(mode="ocr", ocr_language="eng", ocr_dpi=250),
            )
            text = " ".join(p.text for p in Document(output).paragraphs).upper()
            self.assertIn("SCAN", text)
            self.assertIn("2026", text)


if __name__ == "__main__":
    unittest.main()
