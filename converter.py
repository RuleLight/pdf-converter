from __future__ import annotations

import os
import shutil
import sys
import tempfile
import uuid
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from threading import Event
from typing import Callable, Iterable

import fitz  # PyMuPDF
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ProgressCallback = Callable[[int, int, str], None]


class ConversionError(RuntimeError):
    """Raised when a PDF cannot be converted safely."""


class ConversionCancelled(ConversionError):
    """Raised when the user cancels a conversion."""


@dataclass(frozen=True)
class PdfInfo:
    pages: int
    size_bytes: int
    encrypted: bool
    title: str


@dataclass(frozen=True)
class ConversionOptions:
    mode: str = "layout"  # layout, exact, ocr
    pages: str = ""
    password: str = ""
    ocr_language: str = "eng"
    ocr_dpi: int = 220


def _open_pdf(pdf_path: str | os.PathLike[str], password: str = "") -> fitz.Document:
    try:
        pdf = fitz.open(str(pdf_path))
    except Exception as exc:  # PyMuPDF exposes several version-specific errors.
        raise ConversionError(f"无法打开 PDF：{exc}") from exc

    if pdf.needs_pass:
        if not password or not pdf.authenticate(password):
            pdf.close()
            raise ConversionError("PDF 受密码保护，请输入正确密码。")
    return pdf


def inspect_pdf(pdf_path: str | os.PathLike[str], password: str = "") -> PdfInfo:
    path = Path(pdf_path)
    if not path.is_file():
        raise ConversionError("找不到所选 PDF 文件。")
    if path.suffix.lower() != ".pdf":
        raise ConversionError("请选择扩展名为 .pdf 的文件。")

    pdf = _open_pdf(path, password)
    try:
        metadata = pdf.metadata or {}
        return PdfInfo(
            pages=pdf.page_count,
            size_bytes=path.stat().st_size,
            encrypted=bool(pdf.is_encrypted),
            title=(metadata.get("title") or "").strip(),
        )
    finally:
        pdf.close()


def pdf_requires_password(pdf_path: str | os.PathLike[str]) -> bool:
    try:
        pdf = fitz.open(str(pdf_path))
    except Exception as exc:
        raise ConversionError(f"无法打开 PDF：{exc}") from exc
    try:
        return bool(pdf.needs_pass)
    finally:
        pdf.close()


def parse_page_range(value: str, total_pages: int) -> list[int]:
    """Parse a one-based range such as ``1-3, 5`` into zero-based indexes."""
    text = (value or "").strip().lower()
    if not text or text in {"all", "全部"}:
        return list(range(total_pages))

    text = text.replace("，", ",").replace("–", "-").replace("—", "-")
    selected: set[int] = set()
    try:
        for chunk in text.split(","):
            chunk = chunk.strip()
            if not chunk:
                continue
            if "-" in chunk:
                start_text, end_text = (part.strip() for part in chunk.split("-", 1))
                start = int(start_text)
                end = int(end_text)
                if start > end:
                    raise ValueError
                selected.update(range(start - 1, end))
            else:
                selected.add(int(chunk) - 1)
    except ValueError as exc:
        raise ConversionError("页码格式不正确，请使用例如 1-3,5 的格式。") from exc

    if not selected or min(selected) < 0 or max(selected) >= total_pages:
        raise ConversionError(f"页码必须在 1 到 {total_pages} 之间。")
    return sorted(selected)


def _check_cancel(cancel_event: Event | None) -> None:
    if cancel_event and cancel_event.is_set():
        raise ConversionCancelled("转换已取消。")


def _notify(callback: ProgressCallback | None, current: int, total: int, message: str) -> None:
    if callback:
        callback(current, total, message)


def _configure_section(section, width_pt: float, height_pt: float, margin_pt: float = 0.0) -> None:
    section.page_width = Pt(width_pt)
    section.page_height = Pt(height_pt)
    section.top_margin = Pt(margin_pt)
    section.bottom_margin = Pt(margin_pt)
    section.left_margin = Pt(margin_pt)
    section.right_margin = Pt(margin_pt)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)


def _new_document() -> Document:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    return document


def _convert_exact(
    pdf_path: Path,
    output_path: Path,
    page_indexes: Iterable[int],
    password: str,
    callback: ProgressCallback | None,
    cancel_event: Event | None,
) -> None:
    page_indexes = list(page_indexes)
    pdf = _open_pdf(pdf_path, password)
    document = _new_document()
    try:
        for position, page_index in enumerate(page_indexes):
            _check_cancel(cancel_event)
            page = pdf.load_page(page_index)
            width_pt, height_pt = float(page.rect.width), float(page.rect.height)
            section = document.sections[0] if position == 0 else document.add_section(WD_SECTION.NEW_PAGE)
            _configure_section(section, width_pt, height_pt)

            if position > 0:
                # add_section() creates the section-break paragraph on the previous
                # page. Make it as small as possible next to the full-page image.
                break_paragraph = document.paragraphs[-1]
                break_paragraph.paragraph_format.space_before = Pt(0)
                break_paragraph.paragraph_format.space_after = Pt(0)
                break_paragraph.paragraph_format.line_spacing = Pt(1)

            # 144 DPI keeps text sharp while avoiding unreasonably large DOCX files.
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(1)
            run = paragraph.add_run()
            # Leave two points for Word's section-break paragraph to prevent a blank page.
            run.add_picture(BytesIO(pixmap.tobytes("png")), width=Pt(width_pt), height=Pt(max(1, height_pt - 2)))

            _notify(callback, position + 1, len(page_indexes), f"正在还原第 {page_index + 1} 页…")
        document.save(output_path)
    finally:
        pdf.close()


def find_tesseract() -> str | None:
    configured = os.environ.get("TESSERACT_CMD")
    if configured and Path(configured).is_file():
        return configured

    located = shutil.which("tesseract")
    if located:
        return located

    candidates: list[Path] = []
    if sys.platform == "win32":
        for root in (os.environ.get("ProgramFiles"), os.environ.get("LOCALAPPDATA")):
            if root:
                candidates.extend(
                    [
                        Path(root) / "Tesseract-OCR" / "tesseract.exe",
                        Path(root) / "Programs" / "Tesseract-OCR" / "tesseract.exe",
                    ]
                )
    elif sys.platform == "darwin":
        candidates.extend([Path("/opt/homebrew/bin/tesseract"), Path("/usr/local/bin/tesseract")])

    bundle_root = Path(getattr(sys, "_MEIPASS", Path(sys.executable).parent))
    candidates.extend([bundle_root / "tesseract", bundle_root / "tesseract.exe"])
    return str(next((path for path in candidates if path.is_file()), "")) or None


def _convert_ocr(
    pdf_path: Path,
    output_path: Path,
    page_indexes: Iterable[int],
    password: str,
    language: str,
    dpi: int,
    callback: ProgressCallback | None,
    cancel_event: Event | None,
) -> None:
    try:
        import pytesseract
        from PIL import Image
        from pytesseract import Output
    except ImportError as exc:
        raise ConversionError("OCR 组件未安装，请重新运行安装脚本。") from exc

    tesseract_cmd = find_tesseract()
    if not tesseract_cmd:
        raise ConversionError(
            "没有找到 Tesseract OCR。普通 PDF 可改用“保留版式”；扫描件 OCR 的安装方法请查看使用说明。"
        )
    pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    page_indexes = list(page_indexes)
    pdf = _open_pdf(pdf_path, password)
    document = _new_document()
    try:
        for position, page_index in enumerate(page_indexes):
            _check_cancel(cancel_event)
            page = pdf.load_page(page_index)
            width_pt, height_pt = float(page.rect.width), float(page.rect.height)
            section = document.sections[0] if position == 0 else document.add_section(WD_SECTION.NEW_PAGE)
            _configure_section(section, width_pt, height_pt, margin_pt=28)

            scale = dpi / 72.0
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            image = Image.open(BytesIO(pixmap.tobytes("png")))
            try:
                data = pytesseract.image_to_data(image, lang=language, output_type=Output.DICT)
            except pytesseract.TesseractError as exc:
                message = str(exc)
                if "Failed loading language" in message or "Error opening data file" in message:
                    raise ConversionError(f"OCR 缺少语言包“{language}”，请安装后再试。") from exc
                raise ConversionError(f"OCR 识别失败：{message}") from exc

            lines: dict[tuple[int, int, int], list[int]] = {}
            for index, text in enumerate(data.get("text", [])):
                if text and text.strip() and int(float(data["conf"][index])) >= 0:
                    key = (data["block_num"][index], data["par_num"][index], data["line_num"][index])
                    lines.setdefault(key, []).append(index)

            ordered_lines = sorted(lines.values(), key=lambda ids: (data["top"][ids[0]], data["left"][ids[0]]))
            if not ordered_lines:
                paragraph = document.add_paragraph("[此页未识别到文字]")
                paragraph.runs[0].italic = True
            else:
                usable_width_pt = max(1.0, width_pt - 56)
                for ids in ordered_lines:
                    words = [data["text"][i].strip() for i in ids if data["text"][i].strip()]
                    paragraph = document.add_paragraph(" ".join(words))
                    left_pt = max(0.0, data["left"][ids[0]] * 72.0 / dpi - 28)
                    paragraph.paragraph_format.left_indent = Pt(min(left_pt, usable_width_pt * 0.7))
                    paragraph.paragraph_format.space_after = Pt(1.5)
                    heights = [data["height"][i] * 72.0 / dpi for i in ids]
                    estimated_size = min(28.0, max(7.0, sum(heights) / max(1, len(heights)) * 0.78))
                    for run in paragraph.runs:
                        run.font.size = Pt(estimated_size)

            _notify(callback, position + 1, len(page_indexes), f"正在识别第 {page_index + 1} 页…")
        document.save(output_path)
    finally:
        pdf.close()


def _convert_layout(
    pdf_path: Path,
    output_path: Path,
    page_indexes: list[int],
    password: str,
    callback: ProgressCallback | None,
    cancel_event: Event | None,
) -> None:
    try:
        from pdf2docx import Converter
    except ImportError as exc:
        raise ConversionError("版式转换组件未安装，请重新运行安装脚本。") from exc

    _check_cancel(cancel_event)
    _notify(callback, 0, len(page_indexes), "正在分析文字、图片和表格版式…")
    converter = None
    try:
        converter = Converter(str(pdf_path), password=password or None)
        converter.convert(str(output_path), pages=page_indexes)
    except Exception as exc:
        raise ConversionError(f"版式转换失败：{exc}") from exc
    finally:
        if converter is not None:
            converter.close()
    _check_cancel(cancel_event)
    _notify(callback, len(page_indexes), len(page_indexes), "版式转换完成。")


def validate_docx(path: str | os.PathLike[str]) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            required = {"[Content_Types].xml", "word/document.xml"}
            if not required.issubset(names) or archive.testzip() is not None:
                raise ConversionError("生成的 DOCX 文件结构不完整。")
    except (OSError, zipfile.BadZipFile) as exc:
        raise ConversionError("生成的 DOCX 文件无法通过完整性检查。") from exc


def convert_pdf(
    pdf_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str],
    options: ConversionOptions | None = None,
    callback: ProgressCallback | None = None,
    cancel_event: Event | None = None,
) -> Path:
    options = options or ConversionOptions()
    source = Path(pdf_path).expanduser().resolve()
    destination = Path(output_path).expanduser().resolve()

    if source == destination:
        raise ConversionError("输入文件和输出文件不能相同。")
    if destination.suffix.lower() != ".docx":
        destination = destination.with_suffix(".docx")
    destination.parent.mkdir(parents=True, exist_ok=True)

    info = inspect_pdf(source, options.password)
    page_indexes = parse_page_range(options.pages, info.pages)
    if not page_indexes:
        raise ConversionError("没有可转换的页面。")

    temp_name = f".{destination.stem}.{uuid.uuid4().hex}.tmp.docx"
    temp_path = destination.parent / temp_name
    try:
        if options.mode == "layout":
            _convert_layout(source, temp_path, page_indexes, options.password, callback, cancel_event)
        elif options.mode == "exact":
            _convert_exact(source, temp_path, page_indexes, options.password, callback, cancel_event)
        elif options.mode == "ocr":
            _convert_ocr(
                source,
                temp_path,
                page_indexes,
                options.password,
                options.ocr_language,
                options.ocr_dpi,
                callback,
                cancel_event,
            )
        else:
            raise ConversionError(f"不支持的转换模式：{options.mode}")

        validate_docx(temp_path)
        os.replace(temp_path, destination)
        return destination
    finally:
        try:
            temp_path.unlink(missing_ok=True)
        except OSError:
            pass
